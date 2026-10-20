import logging
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch
from xml.etree import ElementTree

from openpyxl import Workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.styles import Border, Font, Side

from bot.excel_parser import (
    is_informational_name,
    parse_people,
    parse_schedule,
    parse_service_events as parse_with_catalog,
    parse_service_rosters,
    resolve_assignments,
)
from bot.services import SERVICES, parse_catalog

# Небольшой каталог для проверки общих правил разбора независимо от
# рабочего reminder_bot_services.json.
TEST_SERVICES = parse_catalog({
    "services": [
        {
            "id": "assembly", "name": "Собрание", "aliases": ["собрание"],
            "columns": [{"headers": ["собрание"], "layout": "pairs"}],
            "tasks": [
                {"name": "Ответственный", "aliases": ["хкц"]},
                {"name": "Ведущий", "aliases": ["ведет"]},
                {"name": "Проповедник",
                 "aliases": ["проповедует", "проповедет"]},
            ],
        },
        {
            "id": "sound", "name": "Звук", "aliases": ["звук"],
            "columns": [{"headers": ["звук"], "layout": "pairs"}],
            "tasks": ["Микшер", "Звук"],
        },
        {
            "id": "kitchen", "name": "Кухня", "aliases": ["кухня"],
            "columns": [{"headers": ["кухня"], "layout": "pairs"}],
            "tasks": ["старший", "в смене", "Пироги"],
        },
        {
            "id": "children", "name": "Детское", "aliases": ["детское"],
            "columns": [{"headers": ["детское"], "layout": "pairs"}],
            "tasks": [],
        },
        {
            "id": "slides", "name": "Слайды", "aliases": ["слайды"],
            "columns": [{"headers": ["слайды"], "layout": "single"}],
            "tasks": [],
        },
    ]
})


def parse_service_events(path, month, logger, services=None):
    return parse_with_catalog(
        path, month, logger, TEST_SERVICES if services is None else services
    )


class ExcelParserTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / "schedule.xlsx"
        self.debug_logger = logging.getLogger("test.excel_parser")

    def _make_workbook(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "октябрь26"
        schedule.append(["СОБРАНИЕ"])
        schedule["A1"].border = Border(left=Side(style="thin"))
        schedule.append(["4 октября"])
        schedule.append(["ХКЦ", "Алексей"])
        schedule.append(["ведет", "Павел"])
        schedule.append(["ПРОПОВЕДУЕТ", "Иван"])
        schedule.append(["11 октября"])
        schedule.append(["ХКЦ", "Другой"])

        people = workbook.create_sheet("списки")
        people.append(["Проповедет"])
        people.append(["В.П."])
        people.append(["Иван", "ivan", 123456789])
        people.append(["ведущий"])
        people.append(["Павел", "pavel", 234567891])
        people.append(["Другая задача"])
        people["A6"].font = Font(bold=True)
        people.append(["Не назначать"])
        people.append(["Евгений"])
        people.append(["ХКЦ"])
        people.append(["Сергей"])
        people.append(["Аркадий"])
        people.append(["Алексей", "alebzor", 181601986])
        workbook.save(self.path)

    def test_is_informational_name_recognizes_markers_and_empty_names(self):
        cases = (
            ("* Задача", (True, "Задача")),
            ("[*] Задача", (True, "Задача")),
            ("[Задача]", (True, "Задача")),
            ("*Задача", (True, "Задача")),
            ("Задача*", (False, "Задача*")),
            ("[]", (True, "")),
            ("*", (True, "")),
            ("Обычная задача", (False, "Обычная задача")),
        )
        for raw, expected in cases:
            with self.subTest(raw=raw):
                self.assertEqual(is_informational_name(raw), expected)

    def test_parse_service_events_keeps_marked_tasks_without_assignees(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append(["СОБРАНИЕ", None, "ЗВУК", None])
        schedule.append([date(2026, 10, 11), None, date(2026, 10, 11), None])
        schedule.append(
            ["* Причастие (В.П.)", "В.П.", "[*] Событие", "Дежурный"]
        )
        schedule.append(["[Праздник]", None, "*Звук"])
        schedule.append(["Ведущий", "Аркадий", None, None])
        workbook.save(self.path)
        logger = MagicMock()

        parsed = parse_service_events(
            self.path, date(2026, 10, 1), logger
        )

        assembly = parsed["assembly"][0].assignments
        self.assertEqual(
            [
                (item.task, item.name, item.is_informational)
                for item in assembly
            ],
            [
                ("Причастие (В.П.)", "В.П.", True),
                ("Праздник", "", True),
                ("Ведущий", "Аркадий", False),
            ],
        )
        sound = parsed["sound"][0].assignments
        self.assertEqual(
            [
                (item.task, item.name, item.is_informational)
                for item in sound
            ],
            [
                ("Событие", "Дежурный", True),
                ("Звук", "", True),
            ],
        )
        logger.debug.assert_any_call(
            "Найдено информационных задач: %s; названия: %s",
            4,
            ["Звук", "Праздник", "Причастие (В.П.)", "Событие"],
        )

    def test_production_catalog_maps_columns_and_tasks(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append([
            "СОБРАНИЕ", None, "ЗВУК", None, "СПЕВКА", "СЛАЙДЫ",
            "ТРАНСЛЯЦИЯ", None, "ОРГ. ПОРЯДКА", None, "КУХНЯ", None,
            "КОФЕ", None, "ДЕТСКОЕ", None, "КНИГИ", "график согласован",
        ])
        sunday = date(2026, 10, 11)
        schedule.append([
            sunday, None, sunday, None, date(2026, 10, 14), sunday,
            sunday, None, sunday, None, sunday, None,
            sunday, None, sunday, None, sunday,
        ])
        schedule.append([
            "ХКЦ", "А", "Звук", "Г", "Е", "Ж", "трансляция", "З",
            "старший", "К", "старший", "Н", None, "П", None, "С", "Р",
        ])
        schedule.append([
            "Проповедует", "Б", "сборка оборудования", "Д", None, None,
            "камера", "И", "в смене", "Л", "Пироги", "О",
        ])
        schedule.append([
            "Вторая часть", "В", None, None, None, None, None, None,
            None, "М",
        ])
        workbook.save(self.path)
        logger = MagicMock()

        parsed = parse_with_catalog(self.path, date(2026, 10, 1), logger)

        def tasks(service_id):
            return [
                (item.task, item.name)
                for event in parsed[service_id]
                for item in event.assignments
            ]

        self.assertEqual(set(parsed), {service.id for service in SERVICES})
        self.assertEqual(
            tasks("assembly"),
            [("Ответственный по ХКЦ", "А"), ("Проповедь", "Б")],
        )
        self.assertEqual(
            parsed["assembly"][0].issues,
            ("«Собрание», 11.10.2026: задача «Вторая часть» не из списка",),
        )
        # Спевка выключена: её колонка (и дата в среду) не читается.
        self.assertEqual(
            [event.event_date for event in parsed["sound"]], [sunday]
        )
        self.assertEqual(tasks("sound"), [("звук", "Г"), ("сборка", "Д")])
        self.assertEqual(
            tasks("broadcast"),
            # Колонки одного служения сводятся в порядке листа.
            [("Слайды", "Ж"), ("Компьютер", "З"), ("Камера", "И")],
        )
        self.assertEqual(
            tasks("order"),
            [("старший", "К"), ("в смене", "Л"), ("в смене", "М")],
        )
        # Кофе выключен; Детское и Книги — служения без задач.
        self.assertEqual(
            tasks("kitchen"), [("старший", "Н"), ("пироги", "О")]
        )
        self.assertEqual(tasks("children"), [("Детское", "С")])
        self.assertEqual(tasks("books"), [("Книги", "Р")])
        self.assertEqual(parsed["office_cleaning"], [])

    def test_disabled_catalog_task_is_skipped_and_unknown_reported(self):
        catalog = parse_catalog({
            "services": [{
                "id": "kitchen", "name": "Кухня", "aliases": [],
                "columns": [{"headers": ["кухня"], "layout": "pairs"}],
                "tasks": ["старший", {"name": "пироги", "enabled": False}],
            }]
        })
        event, _ = self._parse_kitchen_rows(
            [["старший", "А"], ["Пироги", "Б"], ["Посуда", "В"]],
            services=catalog,
        )

        self.assertEqual(
            [(item.task, item.name) for item in event.assignments],
            [("старший", "А")],
        )
        self.assertEqual(
            event.issues,
            ("«Кухня», 11.10.2026: задача «Посуда» не из списка",),
        )

    def _parse_kitchen_rows(self, rows, header="КУХНЯ", services=None):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append([header, None])
        schedule.append([date(2026, 10, 11), None])
        for row in rows:
            schedule.append(row)
        workbook.save(self.path)
        logger = MagicMock()
        parsed = parse_service_events(
            self.path, date(2026, 10, 1), logger, services
        )
        event = next(events[0] for events in parsed.values() if events)
        return event, logger

    def test_rows_without_task_continue_previous_task(self):
        event, _ = self._parse_kitchen_rows([
            ["старший", "А"],
            ["в смене", "Б"],
            [None, "В"],
            [None, "Г"],
            ["Пироги", "Д"],
        ])

        self.assertEqual(
            [(item.task, item.name) for item in event.assignments],
            [
                ("старший", "А"),
                ("в смене", "Б"),
                ("в смене", "В"),
                ("в смене", "Г"),
                ("Пироги", "Д"),
            ],
        )
        self.assertEqual(event.issues, ())

    def test_task_row_without_name_opens_list_for_next_rows(self):
        event, _ = self._parse_kitchen_rows([
            ["старший", None],
            [None, "А"],
            [None, "Б"],
        ])

        self.assertEqual(
            [(item.task, item.name) for item in event.assignments],
            [("старший", "А"), ("старший", "Б")],
        )

    def test_empty_name_cell_ends_task_list_and_reports_orphan_rows(self):
        event, logger = self._parse_kitchen_rows([
            ["старший", "А"],
            [None, None],
            [None, "Б"],
            ["* Причастие", None],
            [None, "В"],
        ])

        self.assertEqual(
            [(item.task, item.name, item.is_informational)
             for item in event.assignments],
            [("старший", "А", False), ("Причастие", "", True)],
        )
        self.assertEqual(
            event.issues,
            ("«Кухня», 11.10.2026: строка без задачи",) * 2,
        )
        logger.warning.assert_any_call(
            "Пропущена строка без задачи: служение=%s, дата=%s",
            "Кухня", date(2026, 10, 11),
        )

    def test_rows_before_first_task_belong_to_service(self):
        event, _ = self._parse_kitchen_rows(
            [[None, "А"], [None, None], [None, "Б"]], header="ДЕТСКОЕ"
        )

        self.assertEqual(
            [(item.task, item.name) for item in event.assignments],
            [("Детское", "А"), ("Детское", "Б")],
        )
        self.assertEqual(event.issues, ())

    def test_untasked_rows_at_block_start_of_tasked_service_are_reported(self):
        event, _ = self._parse_kitchen_rows([
            [None, "А"],
            ["старший", "Б"],
        ])

        self.assertEqual(
            [(item.task, item.name) for item in event.assignments],
            [("старший", "Б")],
        )
        self.assertEqual(
            event.issues, ("«Кухня», 11.10.2026: строка без задачи",)
        )

    def test_fixed_task_column_keeps_names_after_informational_row(self):
        catalog = parse_catalog({
            "services": [{
                "id": "kitchen", "name": "Кухня", "aliases": [],
                "columns": [
                    {"headers": ["кофе"], "layout": "pairs", "task": "кофе"}
                ],
                "tasks": ["кофе"],
            }]
        })
        event, _ = self._parse_kitchen_rows(
            [[None, "А"], ["* Пометка", None], [None, "Б"]],
            header="КОФЕ", services=catalog,
        )

        self.assertEqual(
            [(item.task, item.name) for item in event.assignments],
            [("кофе", "А"), ("Пометка", ""), ("кофе", "Б")],
        )
        self.assertEqual(event.issues, ())

    def test_production_kitchen_rows_before_first_task_are_on_shift(self):
        event, _ = self._parse_kitchen_rows(
            [[None, "А"], ["старший", "Б"]], services=SERVICES
        )

        self.assertEqual(
            [(item.task, item.name) for item in event.assignments],
            [("в смене", "А"), ("старший", "Б")],
        )
        self.assertEqual(event.issues, ())

    def test_comments_row_stops_schedule_parsing(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append(["СОБРАНИЕ", None, "СЛАЙДЫ"])
        schedule.append([date(2026, 10, 11), None, date(2026, 10, 11)])
        schedule.append(["ХКЦ", "Алексей", "Слайды А"])
        schedule.append([None, None, "Комментарии для составления графиков"])
        schedule.append(["Проповедует", "Не из графика", "Не из графика"])
        schedule.append([date(2026, 10, 18), None, date(2026, 10, 18)])
        schedule.append(["ХКЦ", "Не из графика"])
        workbook.save(self.path)

        parsed = parse_service_events(
            self.path, date(2026, 10, 1), self.debug_logger
        )

        self.assertEqual(len(parsed["assembly"]), 1)
        self.assertEqual(
            [(item.task, item.name) for item in parsed["assembly"][0].assignments],
            [("Ответственный", "Алексей")],
        )
        self.assertEqual(
            [item.name for item in parsed["slides"][0].assignments],
            ["Слайды А"],
        )
        self.assertEqual(len(parsed["slides"]), 1)

    def test_comments_row_stops_target_date_schedule(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append(["СОБРАНИЕ"])
        schedule.append([date(2026, 10, 11)])
        schedule.append(["ХКЦ", "Алексей"])
        schedule.append(["КОММЕНТАРИИ ДЛЯ СОСТАВЛЕНИЯ ГРАФИКОВ"])
        schedule.append(["Ведущий", "Не из графика"])
        workbook.save(self.path)

        parsed = parse_schedule(
            self.path, date(2026, 10, 11), self.debug_logger
        )

        self.assertEqual(parsed, {"Ответственный": "Алексей"})

    def test_service_events_support_legacy_schedule_sheet_for_assembly(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "График"
        schedule.append(["Дата", "Задача", "Ответственный"])
        schedule.append([date(2026, 10, 11), None, None])
        schedule.append([None, "ХКЦ", "Алексей"])
        schedule.append([None, "Ведущий", "Павел"])
        schedule.append([date(2026, 10, 18), None, None])
        schedule.append([None, "Проповедует", "Иван"])
        workbook.save(self.path)

        parsed = parse_service_events(
            self.path, date(2026, 10, 1), self.debug_logger
        )

        self.assertEqual(
            [
                (
                    event.event_date,
                    [(item.task, item.name) for item in event.assignments],
                )
                for event in parsed["assembly"]
            ],
            [
                (
                    date(2026, 10, 11),
                    [
                        ("Ответственный", "Алексей"),
                        ("Ведущий", "Павел"),
                    ],
                ),
                (
                    date(2026, 10, 18),
                    [("Проповедник", "Иван")],
                ),
            ],
        )
        self.assertTrue(all(not events for key, events in parsed.items() if key != "assembly"))

    def test_service_events_exclude_dates_outside_requested_month_and_year(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append(["СОБРАНИЕ", None, "ДРУГОЕ"])
        schedule.append([date(2026, 10, 11), None])
        schedule.append(["ХКЦ; Inline value", "Октябрьский ответственный"])
        schedule.append([date(2026, 11, 1), None])
        schedule.append(["ХКЦ", "Ноябрьский ответственный"])
        schedule.append([date(2027, 10, 2), None])
        schedule.append(["ХКЦ", "Следующего года"])
        workbook.save(self.path)

        parsed = parse_service_events(
            self.path, date(2026, 10, 1), self.debug_logger
        )

        self.assertEqual(len(parsed["assembly"]), 1)
        self.assertEqual(
            parsed["assembly"][0].event_date, date(2026, 10, 11)
        )
        self.assertEqual(
            [item.name for item in parsed["assembly"][0].assignments],
            ["Октябрьский ответственный"],
        )

    def test_next_month_date_stops_month_sheet(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append(["СОБРАНИЕ", None, "СЛАЙДЫ"])
        schedule.append([date(2026, 9, 27), None, date(2026, 9, 27)])
        schedule.append(["ХКЦ", "Сентябрьский", "Сентябрьские слайды"])
        schedule.append([date(2026, 10, 25), None, date(2026, 10, 25)])
        schedule.append(["ХКЦ", "Октябрьский", "Октябрьские слайды"])
        schedule.append([None, "Октябрьский второй"])
        schedule.append([None, None, date(2026, 11, 1)])
        schedule.append([None, "Ноябрьский", "Ноябрьские слайды"])
        schedule.append([date(2026, 10, 31), None, date(2026, 10, 31)])
        schedule.append(["ХКЦ", "После ноября", "После ноября"])
        workbook.save(self.path)

        parsed = parse_service_events(
            self.path, date(2026, 10, 1), self.debug_logger
        )

        self.assertEqual(
            [event.event_date for event in parsed["assembly"]],
            [date(2026, 10, 25)],
        )
        self.assertEqual(
            [item.name for item in parsed["assembly"][0].assignments],
            ["Октябрьский", "Октябрьский второй"],
        )
        self.assertEqual(
            [
                [item.name for item in event.assignments]
                for event in parsed["slides"]
            ],
            [["Октябрьские слайды"]],
        )

    def test_december_sheet_stops_at_january_text_date(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ДЕКАБРЬ26"
        schedule.append(["СОБРАНИЕ"])
        schedule.append(["27 декабря"])
        schedule.append(["ХКЦ", "Декабрьский"])
        schedule.append(["3 января"])
        schedule.append(["ХКЦ", "Январский"])
        workbook.save(self.path)

        parsed = parse_service_events(
            self.path, date(2026, 12, 1), self.debug_logger
        )

        self.assertEqual(
            [
                (event.event_date, [item.name for item in event.assignments])
                for event in parsed["assembly"]
            ],
            [(date(2026, 12, 27), ["Декабрьский"])],
        )

    def test_parses_month_sheet_and_schedule_blocks_case_insensitively(self):
        self._make_workbook()

        assignments = parse_schedule(
            self.path, date(2026, 10, 4), self.debug_logger
        )

        self.assertEqual(
            assignments,
            {
                "Ответственный": "Алексей",
                "Ведущий": "Павел",
                "Проповедник": "Иван",
            },
        )

    def test_empty_duplicate_role_does_not_override_or_reject_assignment(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append(["11 октября"])
        schedule.append(["ХКЦ", "Алексей"])
        schedule.append(["ХКЦ", None])
        workbook.save(self.path)

        assignments = parse_schedule(
            self.path, date(2026, 10, 11), self.debug_logger
        )

        self.assertEqual(assignments, {"Ответственный": "Алексей"})

    def test_conflicting_non_empty_duplicate_role_is_rejected(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "ОКТЯБРЬ26"
        schedule.append(["11 октября"])
        schedule.append(["ХКЦ", "Алексей"])
        schedule.append(["ХКЦ", "Павел"])
        workbook.save(self.path)

        with self.assertRaisesRegex(RuntimeError, "указана несколько раз"):
            parse_schedule(self.path, date(2026, 10, 11), self.debug_logger)

    def test_parses_task_sections_anywhere_and_uses_last_filled_person_row(self):
        self._make_workbook()

        people = parse_people(self.path, self.debug_logger)
        assignments = resolve_assignments(
            {
                "Ответственный": "Алексей",
                "Ведущий": "Павел",
                "Проповедник": "Иван",
            },
            people,
        )

        self.assertEqual(assignments["Ответственный"]["tg_id"], 181601986)
        self.assertEqual(assignments["Ответственный"]["username"], "alebzor")
        self.assertEqual(assignments["Проповедник"]["tg_id"], 123456789)
        self.assertNotIn("Не назначать", people["Ведущий"])

    def test_parses_legacy_sheet_names_and_role_labels(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "График"
        schedule.append(["11 октября"])
        schedule.append(["Ответственный", "Алексей"])
        schedule.append(["Проповедник", "Иван"])
        people = workbook.create_sheet("Список")
        people.append(["Ответственный"])
        people.append(["Алексей", "alex", 181601986])
        people.append(["Проповедник"])
        people.append(["Иван", "ivan", 123456789])
        workbook.save(self.path)

        schedule_data = parse_schedule(
            self.path, date(2026, 10, 11), self.debug_logger
        )
        contacts = parse_people(self.path, self.debug_logger)
        assignments = resolve_assignments(schedule_data, contacts)

        self.assertEqual(schedule_data, {
            "Ответственный": "Алексей",
            "Проповедник": "Иван",
        })
        self.assertEqual(assignments["Ответственный"]["tg_id"], 181601986)
        self.assertEqual(assignments["Проповедник"]["tg_id"], 123456789)

    def test_parses_legacy_schedule_inline_and_non_first_column_roles(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "График"
        schedule.append(["11 октября"])
        schedule.append(["ХКЦ: Алексей"])
        schedule.append([None, "Ведущий", "Павел"])
        schedule.append([None, None, "Проповедник; Иван"])
        workbook.save(self.path)

        assignments = parse_schedule(
            self.path, date(2026, 10, 11), self.debug_logger
        )

        self.assertEqual(assignments, {
            "Ответственный": "Алексей",
            "Ведущий": "Павел",
            "Проповедник": "Иван",
        })

    def test_parses_legacy_prefixed_dates_and_space_separated_roles(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "График"
        schedule.append(["Воскресенье, 04.10.2026"])
        schedule.append(["Ведущий Пётр"])
        schedule.append(["Воскресенье, 11.10.2026"])
        schedule.append(["Ведущий Павел"])
        schedule.append(["Ответственный Алексей"])
        schedule.append(["Ведет собрание"])
        schedule.append(["Перенос 18.10.2026"])
        schedule.append(["Воскресенье, 18.10.2026"])
        schedule.append(["Проповедник Иван"])
        workbook.save(self.path)

        assignments = parse_schedule(
            self.path, date(2026, 10, 11), self.debug_logger
        )

        self.assertEqual(assignments, {
            "Ведущий": "Павел",
            "Ответственный": "Алексей",
        })
        events = parse_service_events(
            self.path, date(2026, 10, 1), self.debug_logger
        )["assembly"]
        self.assertEqual(
            [event.event_date for event in events],
            [date(2026, 10, 4), date(2026, 10, 11), date(2026, 10, 18)],
        )

    def test_service_rosters_read_task_blocks_and_contacts(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Списки"
        sheet.append(["Звук", None, None, "Кухня"])
        sheet.append(["звук", "@orphan_user", None, "старший"])
        sheet.append(["Алексей", "@alex_sound", 11111, "Мария"])
        sheet.append(["Перевозка + сборка"])
        sheet.append(["Алексей"])
        sheet.append(["Пётр"])
        sheet.append(["спевка"])
        sheet.append(["Иван"])
        workbook.save(self.path)
        issues = []

        rosters = parse_service_rosters(
            self.path, self.debug_logger, SERVICES, issues
        )

        sound = {member.name: member for member in rosters["sound"]}
        self.assertEqual(sorted(sound), ["Алексей", "Пётр"])
        self.assertEqual(sound["Алексей"].tasks, ("звук", "сборка"))
        self.assertEqual(sound["Алексей"].username, "alex_sound")
        self.assertEqual(sound["Алексей"].telegram_id, 11111)
        self.assertEqual(sound["Пётр"].tasks, ("сборка",))
        self.assertEqual(
            [(member.name, member.tasks) for member in rosters["kitchen"]],
            [("Мария", ("старший",))],
        )
        self.assertNotIn("order", rosters)
        self.assertEqual(
            issues,
            [
                "«Списки», Звук: в строке задачи «звук» указаны "
                "username/Telegram ID без имени"
            ],
        )

    def test_parses_multiple_legacy_assignments_from_one_row(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "График"
        schedule.append(["11 октября"])
        schedule.append([
            "ХКЦ: Алексей",
            "Ведущий; Павел",
            "Проповедует: Иван",
        ])
        workbook.save(self.path)

        assignments = parse_schedule(
            self.path, date(2026, 10, 11), self.debug_logger
        )

        self.assertEqual(assignments, {
            "Ведущий": "Павел",
            "Ответственный": "Алексей",
            "Проповедник": "Иван",
        })

    def test_legacy_role_cell_is_not_used_as_neighboring_role_name(self):
        workbook = Workbook()
        schedule = workbook.active
        schedule.title = "График"
        schedule.append(["11 октября"])
        schedule.append(["Ответственный", "Ведущий: Павел"])
        workbook.save(self.path)

        assignments = parse_schedule(
            self.path, date(2026, 10, 11), self.debug_logger
        )

        self.assertEqual(assignments, {"Ведущий": "Павел"})

    def test_parses_legacy_semicolon_contacts_and_decorated_ids_only(self):
        workbook = Workbook()
        people = workbook.active
        people.title = "Список"
        people.append(["ХКЦ"])
        people.append(["Алексей; @alex; TG_ID: 123456789"])
        people.append(["Мария", "@maria", "Telegram_ID: 234567891"])
        workbook.save(self.path)

        parsed = parse_people(self.path, self.debug_logger)

        self.assertEqual(
            parsed["Ответственный"]["Алексей"],
            {"username": "alex", "tg_id": 123456789},
        )
        self.assertEqual(
            parsed["Ответственный"]["Мария"],
            {"username": "maria", "tg_id": 234567891},
        )

    def test_tabular_contacts_reject_decorated_telegram_ids(self):
        workbook = Workbook()
        people = workbook.active
        people.title = "Списки"
        people.append(["ХКЦ"])
        people.append(["Алексей", "alex", "TG_ID: 123456789"])
        workbook.save(self.path)

        parsed = parse_people(self.path, self.debug_logger)

        self.assertIsNone(parsed["Ответственный"]["Алексей"]["tg_id"])

    def test_partial_resolution_keeps_username_without_telegram_id(self):
        issues = []
        assignments = resolve_assignments(
            {"Ответственный": "Алексей"},
            {
                "Ответственный": {
                    "Алексей": {"username": "alex", "tg_id": None}
                }
            },
            issues=issues,
        )

        self.assertEqual(
            assignments["Ответственный"],
            {"name": "Алексей", "username": "alex", "tg_id": None},
        )
        self.assertTrue(any("отсутствует Telegram_ID" in issue for issue in issues))

    def test_parses_people_when_read_only_cell_has_no_font(self):
        workbook = Workbook()
        people = workbook.active
        people.title = "Списки"
        people.append(["Неизвестная задача"])
        people["A1"].font = Font(bold=True)
        workbook.save(self.path)

        with patch.object(ReadOnlyCell, "font", new=property(lambda _cell: None)):
            parsed_people = parse_people(self.path, self.debug_logger)

        self.assertEqual(parsed_people, {role: {} for role in (
            "Ответственный", "Ведущий", "Проповедник"
        )})

    def test_partial_duplicate_contact_does_not_clear_telegram_id(self):
        workbook = Workbook()
        people = workbook.active
        people.title = "Списки"
        people.append(["ХКЦ"])
        people.append(["Алексей", "alex", 181601986])
        people.append(["Алексей", "alexander", None])
        workbook.save(self.path)

        parsed_people = parse_people(self.path, self.debug_logger)

        self.assertEqual(
            parsed_people["Ответственный"]["Алексей"],
            {"username": "alexander", "tg_id": 181601986},
        )

    def test_selected_person_without_telegram_id_fails_explicitly(self):
        self._make_workbook()

        logger = MagicMock()
        with self.assertRaisesRegex(RuntimeError, "отсутствует Telegram_ID"):
            resolve_assignments(
                {"Ответственный": "Сергей"},
                parse_people(self.path, self.debug_logger),
                logger,
            )

        logger.error.assert_called_once()
        self.assertIn("Сергей", logger.error.call_args.args)
        self.assertIsNone(logger.error.call_args.args[-1])

    def test_logs_resolved_assignment_details(self):
        logger = MagicMock()
        assignments = resolve_assignments(
            {"Ответственный": "Алексей"},
            {
                "Ответственный": {
                    "Алексей": {
                        "username": "alebzor",
                        "tg_id": 181601986,
                    }
                }
            },
            logger,
        )

        self.assertEqual(assignments["Ответственный"]["tg_id"], 181601986)
        logger.debug.assert_called_once()
        self.assertEqual(
            logger.debug.call_args.args[1:],
            ("Ответственный", "Алексей", "alebzor", 181601986),
        )

    def test_repairs_known_invalid_border_style_in_memory(self):
        self._make_workbook()
        malformed_path = Path(self.temp_dir.name) / "malformed.xlsx"
        main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

        with zipfile.ZipFile(self.path) as source:
            with zipfile.ZipFile(malformed_path, "w") as destination:
                for item in source.infolist():
                    content = source.read(item.filename)
                    if item.filename == "xl/styles.xml":
                        root = ElementTree.fromstring(content)
                        borders = root.find(f"{{{main_ns}}}borders")
                        border_side = next(
                            side
                            for border in borders
                            for side in border
                            if side.get("style") == "thin"
                        )
                        border_side.set("style", "solid")
                        content = ElementTree.tostring(
                            root, encoding="utf-8", xml_declaration=True
                        )
                    elif (
                        item.filename.startswith("xl/worksheets/")
                        and item.filename.endswith(".xml")
                    ):
                        root = ElementTree.fromstring(content)
                        dimension = root.find(f"{{{main_ns}}}dimension")
                        if dimension is not None:
                            dimension.set("ref", "A1")
                            content = ElementTree.tostring(
                                root, encoding="utf-8", xml_declaration=True
                            )
                    destination.writestr(item, content)

        schedule = parse_schedule(
            malformed_path, date(2026, 10, 4), self.debug_logger
        )
        people = parse_people(malformed_path, self.debug_logger)
        assignments = resolve_assignments(schedule, people)

        self.assertEqual(assignments["Ведущий"]["name"], "Павел")
        self.assertEqual(assignments["Ответственный"]["tg_id"], 181601986)
        with zipfile.ZipFile(malformed_path) as archive:
            styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
            schedule_xml = ElementTree.fromstring(
                archive.read("xl/worksheets/sheet1.xml")
            )
        self.assertTrue(
            any(side.get("style") == "solid" for side in styles.iter())
        )
        self.assertEqual(
            schedule_xml.find(f"{{{main_ns}}}dimension").get("ref"),
            "A1",
        )


if __name__ == "__main__":
    unittest.main()
