"""Парсинг месячного расписания и справочника участников в xlsx."""
from __future__ import annotations

import logging
import re
import zipfile
from dataclasses import dataclass, replace
from io import BytesIO
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.borders import Side

from .services import (
    SERVICES,
    ScheduleColumn,
    ServiceDefinition,
    normalized_name,
)

ROLES = ("Ответственный", "Ведущий", "Проповедник")


class SheetNotFoundError(RuntimeError):
    """The workbook has no sheet for the requested month."""


class SheetEmptyError(RuntimeError):
    """The month sheet exists but has no rows yet."""


@dataclass(frozen=True)
class ServiceAssignment:
    task: str
    name: str
    is_informational: bool = False


@dataclass(frozen=True)
class ServiceEvent:
    service_id: str
    event_date: date
    assignments: tuple[ServiceAssignment, ...]
    issues: tuple[str, ...] = ()

_MONTH_SHEETS = (
    "ЯНВАРЬ", "ФЕВРАЛЬ", "МАРТ", "АПРЕЛЬ", "МАЙ", "ИЮНЬ",
    "ИЮЛЬ", "АВГУСТ", "СЕНТЯБРЬ", "ОКТЯБРЬ", "НОЯБРЬ", "ДЕКАБРЬ",
)
_MONTH_NAMES = {
    "января": 1, "январь": 1,
    "февраля": 2, "февраль": 2,
    "марта": 3, "март": 3,
    "апреля": 4, "апрель": 4,
    "мая": 5, "май": 5,
    "июня": 6, "июнь": 6,
    "июля": 7, "июль": 7,
    "августа": 8, "август": 8,
    "сентября": 9, "сентябрь": 9,
    "октября": 10, "октябрь": 10,
    "ноября": 11, "ноябрь": 11,
    "декабря": 12, "декабрь": 12,
}
_ROLE_ALIASES = {
    "ответственный": "Ответственный",
    "хкц": "Ответственный",
    "ведущий": "Ведущий",
    "ведет": "Ведущий",
    "проповедует": "Проповедник",
    "проповедет": "Проповедник",
    "проповедник": "Проповедник",
}
_COMMENTS_MARKER = "комментарии"
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _repair_known_stylesheet_issue(path: Path) -> BytesIO | None:
    with zipfile.ZipFile(path) as source:
        try:
            stylesheet = source.read("xl/styles.xml")
        except KeyError:
            return None
        root = ElementTree.fromstring(stylesheet)
        borders = root.find(f"{{{_MAIN_NS}}}borders")
        if borders is None:
            return None

        repaired = False
        valid_styles = Side.style.values
        for border in borders:
            for side in border:
                style = side.get("style")
                if style is None or style in valid_styles:
                    continue
                if style != "solid":
                    raise ValueError(
                        f"Невозможно автоматически исправить стиль границы '{style}'"
                    )
                side.set("style", "thin")
                repaired = True

        if not repaired:
            return None

        repaired_stylesheet = ElementTree.tostring(
            root, encoding="utf-8", xml_declaration=True
        )
        output = BytesIO()
        with zipfile.ZipFile(output, "w") as destination:
            for item in source.infolist():
                data = (
                    repaired_stylesheet
                    if item.filename == "xl/styles.xml"
                    else source.read(item.filename)
                )
                destination.writestr(item, data)
        output.seek(0)
        return output


def _load_workbook(path: Path, debug_logger: logging.Logger):
    try:
        return load_workbook(path, data_only=True, read_only=True)
    except ValueError as exc:
        if "could not read stylesheet" not in str(exc):
            raise
        repaired_workbook = _repair_known_stylesheet_issue(path)
        if repaired_workbook is None:
            raise
        debug_logger.warning(
            "В %s найден некорректный стиль границы 'solid'; "
            "для чтения использован вариант в памяти со стилем 'thin'",
            path,
        )
        return load_workbook(
            repaired_workbook, data_only=True, read_only=True
        )


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_informational_name(raw: str) -> tuple[bool, str]:
    """Return whether a task label is informational and its unmarked name."""
    text = _clean(raw)
    if text.startswith("*"):
        return True, text[1:].strip()
    if text.startswith("["):
        closing = text.find("]")
        if closing >= 0:
            wrapped_name = text[1:closing].strip()
            suffix = text[closing + 1:].strip()
            if wrapped_name == "*" and suffix:
                return True, suffix
            if not suffix:
                return True, wrapped_name
    return False, text


def _normalized(value: str) -> str:
    return " ".join(value.casefold().replace("ё", "е").split())


def _rows_before_comments(rows: list) -> list:
    """Обрезает строки листа, начиная со строки «КОММЕНТАРИИ…»."""
    for index, row in enumerate(rows):
        if any(
            isinstance(value, str)
            and _normalized(value).startswith(_COMMENTS_MARKER)
            for value in row or ()
        ):
            return rows[:index]
    return rows


def _rows_before_other_month(
    rows: list, date_columns, month: date, row_date=None
) -> list:
    """Обрезает лист на первом воскресенье другого месяца после дат месяца.

    Даты предыдущего месяца в начале листа пропускаются, а воскресенье
    следующего месяца завершает чтение: его блок и все строки ниже не
    учитываются. Будние даты (например, спевка в среду после последнего
    воскресенья) лист не обрезают.
    """
    seen_month = False
    for index, row in enumerate(rows):
        for column in date_columns:
            if len(row) <= column:
                continue
            event_date = (row_date or _cell_date)(row[column], month.year)
            if event_date is None:
                continue
            if (event_date.year, event_date.month) == (month.year, month.month):
                seen_month = True
            elif seen_month and event_date.weekday() == 6:
                return rows[:index]
    return rows


def _canonical_role(value: str) -> str | None:
    return _ROLE_ALIASES.get(_normalized(value))


def _schedule_sheet_name(target: date) -> str:
    return f"{_MONTH_SHEETS[target.month - 1]}{target.year % 100:02d}"


def _find_sheet(workbook, expected: str, legacy_names: tuple[str, ...] = ()):
    sheets_by_name = {_normalized(name): name for name in workbook.sheetnames}
    for candidate in (expected, *legacy_names):
        sheet_name = sheets_by_name.get(_normalized(candidate))
        if sheet_name is not None:
            return workbook[sheet_name], sheet_name
    return None, None


def _cell_date(value, target_year: int) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value).casefold().replace("ё", "е")
    numeric = re.match(r"^(\d{1,2})[./](\d{1,2})[./](\d{4})(?:\b|$)", text)
    if numeric:
        try:
            return date(*(int(part) for part in (numeric[3], numeric[2], numeric[1])))
        except ValueError:
            return None
    named = re.match(r"^(\d{1,2})\s+([а-я]+)(?:\s+(\d{4}))?(?:\b|$)", text)
    if named:
        month = _MONTH_NAMES.get(named[2])
        if month is None:
            return None
        try:
            return date(int(named[3] or target_year), month, int(named[1]))
        except ValueError:
            return None
    return None


def _parse_task_row(row) -> tuple[str, str] | None:
    task = _clean(row[0]) if row else ""
    if not task:
        return None
    task_name, separator, inline_name = task.partition(";")
    role = _canonical_role(task_name)
    if role is None:
        return None
    name = _clean(row[1]) if len(row) > 1 else ""
    if not name and separator:
        name = inline_name.strip()
    return role, name


def _legacy_role_cell(value) -> tuple[str, str] | None:
    task = _clean(value)
    if not task:
        return None
    role_text, separator, inline_name = task.partition(";")
    if not separator:
        role_text, separator, inline_name = task.partition(":")
    role = _canonical_role(role_text if separator else task)
    if role is not None:
        return role, inline_name.strip() if separator else ""
    if separator:
        return None
    # «Ведущий Павел»: каноническая роль и имя через пробел, как в прежнем
    # формате; синонимы («Ведет собрание») так не разбираются.
    role_text, _, name = task.partition(" ")
    role = next(
        (role for role in ROLES if _normalized(role) == _normalized(role_text)),
        None,
    )
    if role is None or not name.strip():
        return None
    return role, name.strip()


# Перед датой допускается только день недели: «Воскресенье, 11.10.2026».
_WEEKDAYS = (
    "понедельник", "вторник", "среда", "четверг", "пятница", "суббота",
    "воскресенье", "пн", "вт", "ср", "чт", "пт", "сб", "вс",
)
_LEGACY_DATE_RE = re.compile(
    rf"^(?:{'|'.join(_WEEKDAYS)})\.?,?\s+"
    r"(\d{1,2})[./](\d{1,2})[./](\d{4})(?!\d)",
    re.IGNORECASE,
)


def _legacy_row_date(value, target_year: int) -> date | None:
    """Дата строки листа «График»; допускает текст перед датой."""
    parsed = _cell_date(value, target_year)
    if parsed is not None:
        return parsed
    match = _LEGACY_DATE_RE.match(_clean(value))
    if match is None:
        return None
    try:
        return date(int(match[3]), int(match[2]), int(match[1]))
    except ValueError:
        return None


def _parse_legacy_task_row(row) -> list[tuple[str, str]]:
    assignments = []
    for index, value in enumerate(row):
        parsed_role = _legacy_role_cell(value)
        if parsed_role is None:
            continue
        role, name = parsed_role
        if not name:
            for candidate in row[index + 1:]:
                candidate_text = _clean(candidate)
                if not candidate_text:
                    continue
                if _legacy_role_cell(candidate_text) is not None:
                    break
                name = candidate_text
                break
        if name:
            assignments.append((role, name))
    return assignments


def _parse_legacy_event_row(row) -> list[ServiceAssignment]:
    assignments = [
        ServiceAssignment(task=role, name=name)
        for role, name in _parse_legacy_task_row(row)
    ]
    if assignments:
        return assignments
    raw_task = _clean(row[1]) if len(row) > 1 else ""
    is_informational, task = is_informational_name(raw_task)
    if is_informational and task:
        name = _clean(row[2]) if len(row) > 2 else ""
        assignments.append(
            ServiceAssignment(
                task=task,
                name=name,
                is_informational=True,
            )
        )
    return assignments


def parse_schedule(path: Path, target: date,
                   debug_logger: logging.Logger) -> dict[str, str]:
    """Возвращает назначения трёх ролей для даты target."""
    workbook = _load_workbook(path, debug_logger)
    try:
        expected_sheet = _schedule_sheet_name(target)
        worksheet, sheet_name = _find_sheet(
            workbook, expected_sheet, legacy_names=("График",)
        )
        is_legacy_sheet = sheet_name is not None and _normalized(sheet_name) == "график"
        debug_logger.debug(
            "Открыт %s, листы=%s, лист расписания=%s",
            path, workbook.sheetnames, sheet_name,
        )
        if worksheet is None:
            raise RuntimeError(f"Лист расписания '{expected_sheet}' не найден")

        worksheet.reset_dimensions()
        rows = _rows_before_comments(
            list(worksheet.iter_rows(values_only=True))
        )
        row_date = _legacy_row_date if is_legacy_sheet else _cell_date
        target_row = next(
            (
                index for index, row in enumerate(rows)
                if row and row_date(row[0], target.year) == target
            ),
            None,
        )
        if target_row is None:
            raise RuntimeError(
                f"Дата {target:%d.%m.%Y} не найдена на листе '{sheet_name}'"
            )

        found: dict[str, str] = {}
        for row in rows[target_row + 1:]:
            if row and row_date(row[0], target.year) is not None:
                break
            if is_legacy_sheet:
                task_rows = _parse_legacy_task_row(row)
            else:
                task_row = _parse_task_row(row)
                task_rows = [task_row] if task_row is not None else []
            for role, name in task_rows:
                if name and role in found and found[role] != name:
                    raise RuntimeError(
                        f"Роль «{role}» указана несколько раз на дату {target:%d.%m.%Y}"
                    )
                if name:
                    found[role] = name
                    debug_logger.debug("Найдено: %s = %s", role, name)

        missing = [role for role in ROLES if role not in found]
        if missing:
            debug_logger.error(
                "Не найдены назначения на листе '%s' для %s: %s",
                sheet_name, target.strftime("%d.%m.%Y"), missing,
            )
        debug_logger.debug("Итог парсинга '%s': %s", sheet_name, found)
        return found
    finally:
        workbook.close()


def _column_assignments(
    rows: list,
    row_index: int,
    next_row: int,
    start_column: int,
    column: ScheduleColumn,
    service: ServiceDefinition,
    event_date: date,
    month: date,
    debug_logger: logging.Logger,
) -> tuple[list[ServiceAssignment], list[str]]:
    """Читает назначения одной колонки служения в блоке одной даты."""
    default_task = column.task or service.name
    assignments: list[ServiceAssignment] = []
    issues: list[str] = []
    # Строки без задачи продолжают последнюю задачу блока, пока не встретится
    # пустая ячейка имени, новая задача или пометка. Иначе они относятся к
    # задаче колонки или к служению без задач; в начале блока — к задаче
    # default_task колонки; в остальных случаях это «строка без задачи».
    current_task: str | None = None
    block_has_task = False
    for row in rows[row_index + 1:next_row]:
        if column.layout == "single":
            name = _clean(row[start_column]) if len(row) > start_column else ""
            is_informational, task = is_informational_name(name)
            if is_informational and task:
                assignments.append(
                    ServiceAssignment(task=task, name="", is_informational=True)
                )
            elif name and _cell_date(name, month.year) is None:
                assignments.append(ServiceAssignment(task=default_task, name=name))
            continue

        raw_task_text = (
            _clean(row[start_column]) if len(row) > start_column else ""
        )
        is_informational, task_text = is_informational_name(raw_task_text)
        name = (
            _clean(row[start_column + 1]) if len(row) > start_column + 1 else ""
        )
        task_label, separator, inline_name = task_text.partition(";")
        if not name and separator:
            name = inline_name.strip()
        task_label = task_label.strip()
        if is_informational:
            current_task = None
            block_has_task = True
            task = task_label
        elif task_label:
            task = _canonical_role(task_label) or task_label
            current_task = task
            block_has_task = True
        elif not name:
            current_task = None
            continue
        elif current_task is not None:
            task = current_task
        elif column.task or not service.tasks:
            # Задача колонки («кофе») или служение без задач (Книги)
            # относятся ко всем строкам без своей задачи, и после пометок.
            task = default_task
        elif not block_has_task and column.default_task:
            task = column.default_task
        else:
            issues.append(
                f"«{service.name}», {event_date:%d.%m.%Y}: строка без задачи"
            )
            debug_logger.warning(
                "Пропущена строка без задачи: служение=%s, дата=%s",
                service.name, event_date,
            )
            continue
        if not name and not is_informational:
            continue
        if not task:
            continue
        assignments.append(
            ServiceAssignment(
                task=task, name=name, is_informational=is_informational
            )
        )
    return assignments, issues


def _task_enabled(service: ServiceDefinition, task_name: str) -> bool:
    key = normalized_name(task_name)
    return any(
        task.enabled
        for task in service.tasks
        if key in {normalized_name(name) for name in (task.name, *task.aliases)}
    )


def _catalog_assignments(
    service: ServiceDefinition,
    assignments,
    issues: list[str],
    event_date: date,
    debug_logger: logging.Logger,
) -> tuple[ServiceAssignment, ...]:
    """Приводит задачи к названиям каталога и отбрасывает лишние.

    Выключенные задачи пропускаются молча, задачи не из каталога — с
    предупреждением. У служения без задач все назначения относятся к нему.
    """
    tasks = {
        normalized_name(key): task
        for task in service.tasks
        for key in (task.name, *task.aliases)
    }
    result = []
    for assignment in assignments:
        if assignment.is_informational:
            result.append(assignment)
            continue
        if not service.tasks:
            result.append(replace(assignment, task=service.name))
            continue
        task = tasks.get(normalized_name(assignment.task))
        if task is None:
            issue = (
                f"«{service.name}», {event_date:%d.%m.%Y}: "
                f"задача «{assignment.task}» не из списка"
            )
            if issue not in issues:
                issues.append(issue)
            debug_logger.warning(
                "Задача не из каталога: служение=%s, дата=%s, задача=%s",
                service.name, event_date, assignment.task,
            )
            continue
        if not task.enabled:
            debug_logger.debug(
                "Выключенная задача пропущена: служение=%s, задача=%s",
                service.name, task.name,
            )
            continue
        result.append(replace(assignment, task=task.name))
    return tuple(result)


def parse_service_events(
    path: Path,
    month: date,
    debug_logger: logging.Logger,
    services: tuple[ServiceDefinition, ...] | None = None,
) -> dict[str, list[ServiceEvent]]:
    """Читает события и назначения служений с месячного листа.

    services — каталог с актуальными задачами (по умолчанию из файла).
    """
    services = SERVICES if services is None else services
    services_by_id = {service.id: service for service in services}
    workbook = _load_workbook(path, debug_logger)
    try:
        expected_sheet = _schedule_sheet_name(month)
        worksheet, sheet_name = _find_sheet(
            workbook, expected_sheet, legacy_names=("График",)
        )
        if worksheet is None:
            raise SheetNotFoundError(
                f"Лист расписания '{expected_sheet}' не найден"
            )

        worksheet.reset_dimensions()
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            raise SheetEmptyError(f"Лист расписания '{sheet_name}' пуст")
        sheet_width = max(len(row) for row in rows)
        rows = _rows_before_comments(rows)

        result: dict[str, list[ServiceEvent]] = {
            service.id: [] for service in services
        }
        if _normalized(sheet_name) == "график":
            assembly = services_by_id.get("assembly")
            if assembly is None:
                return result
            rows = _rows_before_other_month(
                rows, (0,), month, _legacy_row_date
            )
            event_rows = [
                (row_index, event_date)
                for row_index, row in enumerate(rows)
                if row
                and (event_date := _legacy_row_date(row[0], month.year)) is not None
            ]
            for event_index, (row_index, event_date) in enumerate(event_rows):
                if (event_date.year, event_date.month) != (month.year, month.month):
                    continue
                next_row = (
                    event_rows[event_index + 1][0]
                    if event_index + 1 < len(event_rows)
                    else len(rows)
                )
                issues: list[str] = []
                assignments = _catalog_assignments(
                    assembly,
                    [
                        assignment
                        for row in rows[row_index + 1:next_row]
                        for assignment in _parse_legacy_event_row(row)
                    ],
                    issues,
                    event_date,
                    debug_logger,
                )
                result["assembly"].append(
                    ServiceEvent(
                        service_id="assembly",
                        event_date=event_date,
                        assignments=assignments,
                        issues=tuple(issues),
                    )
                )
            debug_logger.debug(
                "Прочитан legacy-лист '%s': событий Собрания=%s",
                sheet_name,
                len(result["assembly"]),
            )
            _log_informational_tasks(result, debug_logger)
            return result

        columns_by_header = {
            normalized_name(header): (service, column)
            for service in services
            for column in service.columns
            for header in column.headers
        }
        headers = []
        found_columns: list[tuple[ServiceDefinition, ScheduleColumn, int]] = []
        seen_columns: set[tuple[str, tuple[str, ...]]] = set()
        for column_index, value in enumerate(rows[0]):
            header = _clean(value)
            if not header:
                continue
            headers.append(column_index)
            match = columns_by_header.get(normalized_name(header))
            if match is None:
                debug_logger.debug(
                    "Колонка '%s' листа '%s' не относится к служениям",
                    header, sheet_name,
                )
                continue
            service, column = match
            if column.task and not _task_enabled(service, column.task):
                debug_logger.debug(
                    "Колонка '%s' пропущена: задача «%s» выключена",
                    header, column.task,
                )
                continue
            key = (service.id, column.headers)
            if key in seen_columns:
                raise RuntimeError(
                    f"Заголовок «{header}» указан несколько раз "
                    f"на листе '{sheet_name}'"
                )
            seen_columns.add(key)
            found_columns.append((service, column, column_index))

        rows = _rows_before_other_month(
            rows, tuple(index for _, _, index in found_columns), month
        )
        found_ids = {service.id for service, _, _ in found_columns}
        for service in services:
            if service.columns and service.id not in found_ids:
                debug_logger.warning(
                    "На листе '%s' не найдено служение '%s'",
                    sheet_name, service.name,
                )

        # Несколько колонок одного служения (Трансляция + Слайды) дают одно
        # событие на дату; назначения идут в порядке колонок листа.
        collected: dict[str, dict[date, tuple[list, list]]] = {}
        for service, column, start_column in found_columns:
            next_headers = [index for index in headers if index > start_column]
            end_column = (
                min(next_headers) - 1 if next_headers else sheet_width - 1
            )
            if column.layout == "pairs" and end_column < start_column + 1:
                raise RuntimeError(
                    f"Для служения «{service.name}» на листе '{sheet_name}' "
                    "не хватает столбца ответственного"
                )
            event_rows = [
                (row_index, event_date)
                for row_index, row in enumerate(rows)
                if len(row) > start_column
                and (event_date := _cell_date(row[start_column], month.year))
                is not None
            ]
            for event_index, (row_index, event_date) in enumerate(event_rows):
                next_row = (
                    event_rows[event_index + 1][0]
                    if event_index + 1 < len(event_rows)
                    else len(rows)
                )
                if (event_date.year, event_date.month) != (month.year, month.month):
                    continue
                assignments, issues = _column_assignments(
                    rows, row_index, next_row, start_column, column,
                    service, event_date, month, debug_logger,
                )
                event_assignments, event_issues = collected.setdefault(
                    service.id, {}
                ).setdefault(event_date, ([], []))
                event_assignments.extend(assignments)
                event_issues.extend(issues)

        for service in services:
            for event_date, (assignments, issues) in sorted(
                collected.get(service.id, {}).items()
            ):
                normalized = _catalog_assignments(
                    service, assignments, issues, event_date, debug_logger
                )
                result[service.id].append(
                    ServiceEvent(
                        service_id=service.id,
                        event_date=event_date,
                        assignments=normalized,
                        issues=tuple(issues),
                    )
                )
                debug_logger.debug(
                    "Прочитано служение '%s' на %s: назначений=%s",
                    service.name, event_date, len(normalized),
                )

        _log_informational_tasks(result, debug_logger)
        return result
    finally:
        workbook.close()


def _log_informational_tasks(
    events_by_service: dict[str, list[ServiceEvent]],
    debug_logger: logging.Logger,
) -> None:
    task_names = sorted(
        {
            assignment.task
            for events in events_by_service.values()
            for event in events
            for assignment in event.assignments
            if assignment.is_informational
        }
    )
    debug_logger.debug(
        "Найдено информационных задач: %s; названия: %s",
        len(task_names),
        task_names,
    )


def _telegram_id(value, allow_decorated: bool = False) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        return int(value) if value.is_integer() and value > 0 else None
    text = _clean(value)
    if allow_decorated:
        decorated = re.fullmatch(
            r"(?:TG[_ ]?ID|TELEGRAM[_ ]?ID)\s*[:=#-]\s*(\d+)",
            text,
            flags=re.IGNORECASE,
        )
        if decorated:
            text = decorated.group(1)
    return int(text) if text.isdigit() and int(text) > 0 else None


def parse_people(path: Path,
                 debug_logger: logging.Logger) -> dict[str, dict[str, dict]]:
    """Возвращает справочник участников по задаче/роли с листа «Списки»."""
    workbook = _load_workbook(path, debug_logger)
    try:
        worksheet, sheet_name = _find_sheet(
            workbook, "Списки", legacy_names=("Список",)
        )
        if worksheet is None:
            raise RuntimeError("Лист 'Списки' не найден")
        is_legacy_sheet = _normalized(sheet_name) == "список"

        worksheet.reset_dimensions()
        people: dict[str, dict[str, dict]] = {role: {} for role in ROLES}
        current_role: str | None = None
        for raw_row in worksheet.iter_rows():
            cells = list(raw_row)
            first_cell = cells[0] if cells else None
            first = _clean(first_cell.value) if first_cell else ""
            role = _canonical_role(first)
            if role is not None:
                current_role = role
                debug_logger.debug("Лист '%s': начат блок '%s'", sheet_name, first)
                continue
            has_extra_values = any(
                _clean(cell.value) for cell in cells[1:3]
            )
            if (
                first
                and not has_extra_values
                and first_cell is not None
                and getattr(getattr(first_cell, "font", None), "bold", False)
            ):
                current_role = None
                debug_logger.debug(
                    "Лист '%s': пропущен блок неизвестной задачи '%s'",
                    sheet_name, first,
                )
                continue
            if current_role is None or not first:
                continue

            raw_values = [
                _clean(cell.value) if cell is not None else ""
                for cell in cells[:3]
            ]
            if is_legacy_sheet and any(";" in value for value in raw_values):
                legacy_values = [
                    part.strip()
                    for value in raw_values
                    for part in value.split(";")
                    if part.strip()
                ]
                name = legacy_values[0] if legacy_values else first
                id_value = next(
                    (
                        value for value in legacy_values
                        if _telegram_id(value, allow_decorated=True) is not None
                    ),
                    None,
                )
                username = next(
                    (
                        value for value in legacy_values[1:]
                        if _telegram_id(value, allow_decorated=True) is None
                    ),
                    "",
                )
                tg_id = _telegram_id(id_value, allow_decorated=True)
            else:
                name = first
                username = raw_values[1] if len(raw_values) > 1 else ""
                tg_id = _telegram_id(
                    raw_values[2] if len(raw_values) > 2 else None,
                    allow_decorated=is_legacy_sheet,
                )
            existing = people[current_role].get(name, {})
            people[current_role][name] = {
                "username": username.lstrip("@").strip()
                or existing.get("username", ""),
                "tg_id": (
                    tg_id if tg_id is not None else existing.get("tg_id")
                ),
            }

        debug_logger.debug(
            "Лист '%s' загружен: %s",
            sheet_name, {role: list(names) for role, names in people.items()},
        )
        return people
    finally:
        workbook.close()


@dataclass(frozen=True)
class RosterMember:
    """Человек из состава служения на листе «Списки»."""
    name: str
    tasks: tuple[str, ...] = ()
    username: str = ""
    telegram_id: int | None = None


_USERNAME_CELL = re.compile(r"@?([A-Za-z][A-Za-z0-9_]{4,31})")


def parse_service_rosters(
    path: Path,
    debug_logger: logging.Logger,
    services: tuple[ServiceDefinition, ...] | None = None,
    issues: list[str] | None = None,
) -> dict[str, list[RosterMember]]:
    """Читает состав служений с листа «Списки».

    В строке 1 — названия служений; блок служения тянется вправо до
    следующего названия. Ячейка с названием или синонимом задачи служения
    начинает задачу в своей колонке; остальные ячейки ниже — люди этой
    задачи. Люди под выключенной задачей пропускаются. Справа от имени
    могут стоять username и Telegram ID (как в блоке Собрания).
    В результате только служения, у которых есть блок на листе.
    """
    services = SERVICES if services is None else services
    issues = [] if issues is None else issues
    workbook = _load_workbook(path, debug_logger)
    try:
        worksheet, sheet_name = _find_sheet(workbook, "Списки")
        if worksheet is None:
            raise RuntimeError("Лист 'Списки' не найден")
        worksheet.reset_dimensions()
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        if not rows:
            raise RuntimeError("Лист 'Списки' пуст")
        aliases = {
            normalized_name(alias): service
            for service in services
            for alias in (service.name, *service.aliases)
        }
        headers: list[tuple[int, ServiceDefinition | None]] = []
        for column, value in enumerate(rows[0]):
            header = _clean(value)
            if not header:
                continue
            service = aliases.get(normalized_name(header))
            if service is None:
                issues.append(f"«Списки»: служение «{header}» не найдено в каталоге")
                debug_logger.warning(
                    "Лист 'Списки': неизвестное служение '%s'", header
                )
            headers.append((column, service))
        width = max(len(row) for row in rows)
        # Служения без блока на листе в результат не попадают: их состав
        # не трогается (иначе опечатка в заголовке сняла бы всех из списка).
        result: dict[str, list[RosterMember]] = {}
        for index, (start, service) in enumerate(headers):
            if service is None:
                continue
            end = headers[index + 1][0] if index + 1 < len(headers) else width
            task_by_key = {
                normalized_name(key): task
                for task in service.tasks
                for key in (task.name, *task.aliases)
            }
            people: dict[str, dict] = {}
            consumed: set[tuple[int, int]] = set()
            for column in range(start, end):
                current_task = None
                for row_index in range(1, len(rows)):
                    row = rows[row_index]
                    if (row_index, column) in consumed or column >= len(row):
                        continue
                    text = _clean(row[column])
                    if not text:
                        # Username/ID рядом с пустым именем — не люди. Без
                        # имени username узнаём только по «@» или по ID
                        # правее (латинское имя или задача — не username).
                        def cell(offset):
                            neighbour = column + offset
                            if (
                                neighbour < end
                                and neighbour < len(row)
                                and (row_index, neighbour) not in consumed
                            ):
                                return _clean(row[neighbour])
                            return ""
                        has_id = _telegram_id(cell(2)) is not None if cell(2) else False
                        username_cell = cell(1)
                        has_username = bool(
                            username_cell
                            and _USERNAME_CELL.fullmatch(username_cell)
                            and normalized_name(username_cell) not in task_by_key
                            and (username_cell.startswith("@") or has_id)
                        )
                        if has_username:
                            consumed.add((row_index, column + 1))
                        if has_id:
                            consumed.add((row_index, column + 2))
                        orphan = has_username or has_id
                        if orphan:
                            issues.append(
                                f"«Списки», {service.name}: строка "
                                f"{row_index + 1} — username/Telegram ID без "
                                "имени"
                            )
                        continue
                    task = task_by_key.get(normalized_name(text))
                    if task is not None:
                        current_task = task
                        # Username и Telegram ID в строке с названием задачи
                        # не принадлежат никому — сообщаем, но не читаем их
                        # как людей.
                        orphan = False
                        for offset, looks_right in (
                            (1, lambda v: _USERNAME_CELL.fullmatch(_clean(v))),
                            (2, lambda v: _telegram_id(v) is not None),
                        ):
                            neighbour = column + offset
                            if (
                                neighbour < end
                                and neighbour < len(row)
                                and looks_right(row[neighbour])
                            ):
                                consumed.add((row_index, neighbour))
                                orphan = True
                        if orphan:
                            issues.append(
                                f"«Списки», {service.name}: в строке задачи "
                                f"«{text}» указаны username/Telegram ID без "
                                "имени"
                            )
                        continue
                    username = ""
                    telegram_id = None
                    if column + 1 < end and column + 1 < len(row):
                        match = _USERNAME_CELL.fullmatch(_clean(row[column + 1]))
                        if match:
                            username = match.group(1)
                            consumed.add((row_index, column + 1))
                    if column + 2 < end and column + 2 < len(row):
                        telegram_id = _telegram_id(row[column + 2])
                        if telegram_id is not None:
                            consumed.add((row_index, column + 2))
                    if service.tasks and current_task is None:
                        issues.append(
                            f"«Списки», {service.name}: «{text}» указан "
                            "без задачи"
                        )
                        continue
                    if current_task is not None and not current_task.enabled:
                        continue
                    person = people.setdefault(
                        text, {"tasks": [], "username": "", "telegram_id": None}
                    )
                    if current_task is not None and (
                        current_task.name not in person["tasks"]
                    ):
                        person["tasks"].append(current_task.name)
                    person["username"] = person["username"] or username
                    person["telegram_id"] = person["telegram_id"] or telegram_id
            result[service.id] = [
                RosterMember(
                    name=name,
                    tasks=tuple(info["tasks"]),
                    username=info["username"],
                    telegram_id=info["telegram_id"],
                )
                for name, info in people.items()
            ]
            debug_logger.debug(
                "Лист 'Списки': служение %s, людей=%s",
                service.id, len(result[service.id]),
            )
        return result
    finally:
        workbook.close()


def resolve_assignments(schedule: dict[str, str],
                        people: dict[str, dict[str, dict]],
                        debug_logger: logging.Logger | None = None,
                        issues: list[str] | None = None,
                        ) -> dict[str, dict]:
    """Сопоставляет назначения с Telegram-данными по роли."""
    result: dict[str, dict] = {}
    for role, name in schedule.items():
        bucket = people.get(role, {})
        if name not in bucket:
            message = f"Имя '{name}' (роль «{role}») не найдено на листе 'Списки'"
            if debug_logger is not None:
                debug_logger.error(
                    "Не найден участник: роль=%s, имя=%s; в соответствующем "
                    "блоке справочника такая запись отсутствует",
                    role, name,
                )
            if issues is not None:
                issues.append(message)
                result[role] = {"name": name, "username": "", "tg_id": None}
                continue
            raise RuntimeError(message)
        info = bucket[name]
        if not info.get("tg_id"):
            message = f"У '{name}' (роль «{role}») отсутствует Telegram_ID"
            if debug_logger is not None:
                debug_logger.error(
                    "Найдена запись без Telegram ID: роль=%s, имя=%s, "
                    "username=%s, tg_id=%s",
                    role, name, info.get("username", ""), info.get("tg_id"),
                )
            if issues is not None:
                issues.append(message)
                result[role] = {
                    "name": name,
                    "username": info.get("username", ""),
                    "tg_id": None,
                }
                continue
            raise RuntimeError(message)
        result[role] = {
            "name": name,
            "username": info["username"],
            "tg_id": info["tg_id"],
        }
        if debug_logger is not None:
            debug_logger.debug(
                "Назначение сопоставлено: роль=%s, имя=%s, username=@%s, tg_id=%s",
                role, name, info["username"], info["tg_id"],
            )
    return result
